# -*- coding: utf-8 -*-
# Copyright (c) 2025 relakkes@gmail.com
#
# This file is part of MediaCrawler project.
# Repository: https://github.com/NanmiCoder/MediaCrawler
# GitHub: https://github.com/NanmiCoder
# Licensed under NON-COMMERCIAL LEARNING LICENSE 1.1
#

# 声明：本代码仅供学习和研究目的使用。使用者应遵守以下原则：
# 1. 不得用于任何商业用途。
# 2. 使用时应遵守目标平台的使用条款和robots.txt规则。
# 3. 不得进行大规模爬取或对平台造成运营干扰。
# 4. 应合理控制请求频率，避免给目标平台带来不必要的负担。
# 5. 不得用于任何非法或不当的用途。
#
# 详细许可条款请参阅项目根目录下的LICENSE文件。
# 使用本代码即表示您同意遵守上述原则和LICENSE中的所有条款。

import json
import pathlib
from typing import Dict, List
import aiofiles

from model.m_douban import DoubanPost, DoubanComment
from tools import utils
from var import source_keyword_var, crawler_type_var
import config

from .douban_store_media import DoubanStoreImage


# In-memory storage for posts (with deduplication)
_posts_data: List[Dict] = []
_saved_post_ids: set = set()  # Track saved post IDs to prevent duplicates


class DoubanStoreFactory:
    """Factory for creating Douban store instances"""
    
    @staticmethod
    async def save_post(post: DoubanPost) -> None:
        """Save post to in-memory storage with deduplication"""
        # Check if post already exists
        if post.post_id in _saved_post_ids:
            utils.logger.info(f"[DoubanStoreFactory.save_post] Post {post.post_id} already saved, skipping duplicate")
            return
        
        post_dict = post.dict()
        _posts_data.append(post_dict)
        _saved_post_ids.add(post.post_id)
        utils.logger.info(f"[DoubanStoreFactory.save_post] Saved post {post.post_id} (Total: {len(_posts_data)})")

    @staticmethod
    async def save_data_to_file() -> None:
        """Save all posts to file"""
        if not _posts_data:
            utils.logger.info("[DoubanStoreFactory.save_data_to_file] No posts to save")
            return
        
        if config.SAVE_DATA_OPTION == "json":
            await DoubanStoreFactory._save_to_json()
        elif config.SAVE_DATA_OPTION == "excel":
            await DoubanStoreFactory._save_to_excel()
        # Support other formats as needed

    @staticmethod
    async def _save_to_json() -> None:
        """Save posts to JSON file"""
        file_path = f"data/douban/{source_keyword_var.get()}_posts.json"
        pathlib.Path(file_path).parent.mkdir(parents=True, exist_ok=True)
        
        async with aiofiles.open(file_path, 'w', encoding='utf-8') as f:
            await f.write(json.dumps(_posts_data, ensure_ascii=False, indent=2))
            utils.logger.info(f"[DoubanStoreFactory._save_to_json] Saved {len(_posts_data)} posts to {file_path}")

    @staticmethod
    async def _save_to_excel() -> None:
        """Save posts to Excel file"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            utils.logger.error("[DoubanStoreFactory._save_to_excel] openpyxl not installed")
            return
        
        file_path = f"data/douban/{source_keyword_var.get()}_posts.xlsx"
        pathlib.Path(file_path).parent.mkdir(parents=True, exist_ok=True)
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Posts"
        
        # Write headers
        headers = ["Post ID", "Title", "Content", "Author", "Created Time", "Likes", "Replies", "Images"]
        ws.append(headers)
        
        # Style headers
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        
        # Write data
        for post in _posts_data:
            # 内容可能很长，保存时不截断
            content = post.get("content", "")
            # 但在 Excel 中为了显示效果，如果内容太长可以截断显示，但保存完整内容
            ws.append([
                post.get("post_id", ""),
                post.get("title", ""),
                content,  # 保存完整内容
                post.get("user_nickname", ""),
                post.get("created_time", ""),
                post.get("like_count", 0),
                post.get("reply_count", 0),
                len(post.get("image_urls", []))
            ])
        
        wb.save(file_path)
        utils.logger.info(f"[DoubanStoreFactory._save_to_excel] Saved {len(_posts_data)} posts to {file_path}")


async def save_post(post: DoubanPost) -> None:
    """Save post"""
    await DoubanStoreFactory.save_post(post)


async def save_data_to_file() -> None:
    """Save data to file"""
    await DoubanStoreFactory.save_data_to_file()


async def update_douban_post_image(post_id: str, index: int, pic_content: bytes, image_url: str) -> None:
    """
    Save post image
    """
    # Extract file extension from URL
    extension = image_url.split(".")[-1] if "." in image_url else "jpg"
    extension = extension.split("?")[0]  # Remove query params
    
    pic_id = f"{post_id}_{index}"
    await DoubanStoreImage().store_image({
        "pic_id": pic_id,
        "pic_content": pic_content,
        "extension_file_name": extension
    })
