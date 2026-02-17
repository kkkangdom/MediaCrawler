#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Quick JSON to Excel converter for MediaCrawler
Directly converts existing JSON files to Excel format
"""

import json
import sys
from pathlib import Path
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Error: openpyxl is required. Install it with: pip install openpyxl")
    sys.exit(1)


def export_json_to_excel(json_file: str, output_file: str = None) -> str:
    """
    Convert JSON file to Excel
    
    Args:
        json_file: Path to JSON file
        output_file: Output Excel file path (optional)
        
    Returns:
        Path to generated Excel file
    """
    json_path = Path(json_file)
    
    if not json_path.exists():
        print(f"❌ Error: JSON file not found: {json_file}")
        sys.exit(1)
    
    print(f"📖 Loading JSON from: {json_file}")
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    if not isinstance(data, list):
        print("❌ Error: JSON data must be a list of items")
        sys.exit(1)
    
    if len(data) == 0:
        print("❌ Error: JSON data is empty")
        sys.exit(1)
    
    print(f"✅ Loaded {len(data)} items")
    
    # Create Excel workbook
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet("Contents")
    
    # Get all unique keys
    headers = set()
    for item in data:
        if isinstance(item, dict):
            headers.update(item.keys())
    
    headers = sorted(list(headers))
    print(f"📊 Found {len(headers)} columns")
    
    # Write headers with styling
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for col_num, header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = header_border
    
    # Write data rows
    data_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for row_num, item in enumerate(data, 2):
        for col_num, header in enumerate(headers, 1):
            value = item.get(header, "")
            
            # Handle complex types
            if isinstance(value, (list, dict)):
                value = str(value)
            elif value is None:
                value = ""
            
            cell = sheet.cell(row=row_num, column=col_num, value=value)
            cell.border = data_border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        
        if row_num % 100 == 0:
            print(f"  ⏳ Processed {row_num - 1} rows...")
    
    # Auto-adjust column widths
    print("🔧 Adjusting column widths...")
    for column in sheet.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        
        for cell in column:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except (TypeError, AttributeError):
                pass
        
        adjusted_width = min(max(max_length + 2, 10), 50)
        sheet.column_dimensions[column_letter].width = adjusted_width
    
    # Generate output filename
    if output_file is None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_file = f"{json_path.stem}_{timestamp}.xlsx"
    
    output_path = json_path.parent / output_file if not Path(output_file).is_absolute() else Path(output_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    
    # Save workbook
    print(f"💾 Saving Excel to: {output_path}")
    workbook.save(str(output_path))
    
    print(f"✨ Successfully exported {len(data)} items to: {output_path}")
    return str(output_path)


if __name__ == "__main__":
    import argparse
    
    parser = argparse.ArgumentParser(description="Quick JSON to Excel converter")
    parser.add_argument("json_file", help="Path to JSON file")
    parser.add_argument("-o", "--output", help="Output Excel file path")
    
    args = parser.parse_args()
    
    try:
        export_json_to_excel(args.json_file, args.output)
    except Exception as e:
        print(f"❌ Error: {e}")
        sys.exit(1)
